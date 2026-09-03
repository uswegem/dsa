import openpyxl

from app.services.ingest_branch_manager import parse_branch_manager_file


def _write_workbook(path, headers, rows, sheet_name=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    if sheet_name:
        ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


HEADERS = [
    "DATE", "CLIENT NAME", "CLIENT_CHECK_NO", "CLIENT_ACCOUNT_NO",
    "APPLICATION NUMBER / ESS", "REPORTED_AMOUNT", "BRANCH",
    "DSA  CODE   ",  # deliberately messy whitespace, per spec
    "DSA_ACCOUNT_NO", "DSA_NAME", "DTL_NAME", "DTL_CODE",
]


def test_parses_valid_rows_and_trims_messy_headers(tmp_path):
    path = tmp_path / "branch.xlsx"
    _write_workbook(
        path, HEADERS,
        [
            ["2026-08-05", "Jane Doe", 871, "ACC001", "NGORONGORO", 500000, "Arusha", "DSA01", "AC-DSA01", "Amos", "Grace T", "DTL01"],
            ["2026-08-06", "John Roe", "NCAA1680", "ACC002", "TANAPA", 300000, "Arusha", "DSA02", "AC-DSA02", "Beatrice", "Grace T", "DTL01"],
        ],
    )
    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)
    assert len(result.rows) == 2
    assert not any(i.severity == "ERROR" for i in result.issues)

    r0 = result.rows[0]
    assert r0.client_check_no == "871"  # numeric input, always ingested as text
    r1 = result.rows[1]
    assert r1.client_check_no == "NCAA1680"  # alphanumeric input preserved
    assert r0.dsa_code == "DSA01"  # header whitespace didn't break column resolution


def test_filters_trailing_blank_rows(tmp_path):
    path = tmp_path / "branch.xlsx"
    rows = [["2026-08-05", "Jane Doe", "871", "ACC001", "X", 500000, "Arusha", "DSA01", "AC1", "Amos", "Grace", "DTL01"]]
    rows += [[None] * 12 for _ in range(58)]  # trailing blanks, like the observed 90-row/32-populated sheet
    _write_workbook(path, HEADERS, rows)
    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)
    assert len(result.rows) == 1


def test_date_outside_selected_period_is_rejected_not_silently_accepted(tmp_path):
    path = tmp_path / "branch.xlsx"
    rows = [["2026-08-0" + str(i), f"Client {i}", str(100 + i), f"ACC{i:03d}", "X", 1000, "Arusha",
              f"DSA{i:02d}", f"AC{i}", f"DSA Name {i}", "Grace", "DTL01"] for i in range(1, 6)]
    rows.append(["2026-01-13", "Stray Client", "999", "ACC999", "X", 1000, "Arusha", "DSA99", "AC99", "Stray DSA", "Grace", "DTL01"])
    _write_workbook(path, HEADERS, rows)
    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)

    errors = [i for i in result.issues if i.severity == "ERROR"]
    assert len(errors) == 1
    assert "outside the selected period" in errors[0].message
    assert all(r.dsa_code != "DSA99" for r in result.rows)  # excluded from ingestion, not just flagged
    assert len(result.rows) == 5


def test_missing_or_unparseable_date_is_rejected(tmp_path):
    path = tmp_path / "branch.xlsx"
    _write_workbook(
        path, HEADERS,
        [["", "Jane Doe", "871", "ACC001", "X", 500000, "Arusha", "DSA01", "AC1", "Amos", "Grace", "DTL01"]],
    )
    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)
    assert len(result.rows) == 0
    assert any("DATE is missing or could not be parsed" in i.message for i in result.issues)


def test_missing_required_field_is_an_error_not_a_crash(tmp_path):
    path = tmp_path / "branch.xlsx"
    _write_workbook(
        path, HEADERS,
        [["2026-08-05", "Jane Doe", "871", None, "X", 500000, "Arusha", "DSA01", "AC1", "Amos", "Grace", "DTL01"]],
    )
    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)
    assert len(result.rows) == 0
    assert any("CLIENT_ACCOUNT_NO" in i.message for i in result.issues if i.severity == "ERROR")


def test_multi_sheet_workbook_tags_each_row_with_its_sheet(tmp_path):
    path = tmp_path / "branch_multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Arusha"
    ws1.append(HEADERS)
    ws1.append(["2026-08-05", "Jane Doe", "871", "ACC001", "X", 500000, "Arusha", "DSA01", "AC1", "Amos", "Grace", "DTL01"])
    ws2 = wb.create_sheet("Dodoma")
    ws2.append(HEADERS)
    ws2.append(["2026-08-06", "John Roe", "872", "ACC002", "X", 300000, "Dodoma", "DSA02", "AC2", "Beatrice", "Fred", "DTL02"])
    wb.save(path)

    result = parse_branch_manager_file(str(path), period_month=8, period_year=2026)
    assert len(result.rows) == 2
    sheets = {r.sheet_name for r in result.rows}
    assert sheets == {"Arusha", "Dodoma"}


def test_single_sheet_with_no_matching_columns_raises(tmp_path):
    path = tmp_path / "not_a_roster.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Some", "Random", "Columns"])
    ws.append([1, 2, 3])
    wb.save(path)

    import pytest
    with pytest.raises(ValueError):
        parse_branch_manager_file(str(path), period_month=8, period_year=2026)
