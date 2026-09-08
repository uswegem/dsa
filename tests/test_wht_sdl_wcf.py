"""
DSA-only WHT deduction + SDL/WCF statutory reporting figures (DTL commission
and the DTL Summary report are explicitly untouched by this change - see the
guard tests at the bottom).
"""
import datetime as dt

import openpyxl
import pytest

from app.core.config import get_settings
from app.models import (
    Branch,
    BranchSale,
    BusinessTransaction,
    CommissionAdjustment,
    CommissionLine,
    CommissionRun,
    Dsa,
    Dtl,
    MatchedTransaction,
    Role,
    Upload,
    User,
)
from app.models.enums import LoanType, PayeeType, RunStatus, RunType, UploadStatus, UploadType
from app.reports.excel import SDL_WCF_FOOTNOTE, aggregate_dsa_summary, generate_commission_report
from app.services.commission import calculate_commission_run
from app.services.matching import run_matching

settings = get_settings()


def _seed_base(db):
    branch = Branch(name="Arusha", code="ARU")
    db.add(branch)
    db.flush()

    dsa = Dsa(dsa_code="DSA01", dsa_name="Amos", dsa_account_no="AC-DSA01", branch_id=branch.id)
    dtl = Dtl(dtl_code="DTL01", dtl_name="Grace T", dtl_account_no="AC-DTL01")
    db.add_all([dsa, dtl])

    admin_role = db.query(Role).filter(Role.name == "ADMIN").one()
    admin = User(email="admin@example.com", hashed_password="x", full_name="Admin", role_id=admin_role.id)
    db.add(admin)
    db.flush()

    branch_upload = Upload(
        upload_type=UploadType.BRANCH_MANAGER, uploaded_by_user_id=admin.id, branch_id=branch.id,
        period_month=8, period_year=2026, original_filename="branch.xlsx", storage_path="x", status=UploadStatus.PROCESSED,
    )
    biz_upload = Upload(
        upload_type=UploadType.BUSINESS_MANAGER, uploaded_by_user_id=admin.id, branch_id=None,
        period_month=8, period_year=2026, original_filename="biz.xlsx", storage_path="x", status=UploadStatus.PROCESSED,
    )
    db.add_all([branch_upload, biz_upload])
    db.flush()
    return branch, dsa, dtl, admin, branch_upload, biz_upload


def _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000):
    db.add(BranchSale(
        upload_id=branch_upload.id, branch_id=branch.id, row_number=2, period_month=8, period_year=2026,
        loan_date=dt.date(2026, 8, 5), client_name="Client A", client_check_no="871", client_account_no="ACC001",
        dsa_code="DSA01", dsa_name="Amos", dtl_code="DTL01", dtl_name="Grace T",
    ))
    db.add(BusinessTransaction(
        upload_id=biz_upload.id, row_number=2, branch_name_raw="Arusha", employee_no="871",
        loan_type=LoanType.NL, disbursement_date=dt.date(2026, 8, 5), disbursement_amt=disbursement_amt,
        client_disb_ext_account_no="ACC001",
    ))
    db.flush()
    run_matching(db, branch_upload_id=branch_upload.id, business_upload_id=biz_upload.id)


def _calculate(db, admin, period="2026-08"):
    run = CommissionRun(run_type=RunType.ORG_WIDE, branch_id=None, period=period, status=RunStatus.DRAFT, created_by_user_id=admin.id)
    db.add(run)
    db.flush()
    calculate_commission_run(db, run)
    return run


# ---------------------------------------------------------------------
# Per-line WHT: DSA lines get it, DTL lines are untouched
# ---------------------------------------------------------------------

def test_dsa_line_gets_wht_and_net_commission_amount(db_session):
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    dsa_line = db.query(CommissionLine).filter(CommissionLine.commission_run_id == run.id, CommissionLine.payee_type == PayeeType.DSA).one()
    expected_gross = round(1_000_000 * settings.dsa_nl_rate, 2)  # 70,000
    expected_wht = round(expected_gross * settings.dsa_wht_rate, 2)  # 3,500
    assert float(dsa_line.commission_amount) == expected_gross
    assert float(dsa_line.wht_amount) == expected_wht
    assert float(dsa_line.net_commission_amount) == round(expected_gross - expected_wht, 2)


def test_dtl_line_is_not_wht_deducted(db_session):
    """Scope guard: DTL commission is untouched by this change."""
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    dtl_line = db.query(CommissionLine).filter(CommissionLine.commission_run_id == run.id, CommissionLine.payee_type == PayeeType.DTL).one()
    expected_gross = round(1_000_000 * settings.dtl_nl_rate, 2)
    assert float(dtl_line.commission_amount) == expected_gross  # DTL math itself is unchanged
    assert dtl_line.wht_amount is None
    assert dtl_line.net_commission_amount is None


# ---------------------------------------------------------------------
# aggregate_dsa_summary: gross - wht == net_salary, by construction
# ---------------------------------------------------------------------

def test_aggregate_dsa_summary_wht_and_net_salary_reconcile_with_gross(db_session):
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    dsa_lines = db.query(CommissionLine).filter(CommissionLine.commission_run_id == run.id, CommissionLine.dsa_id.isnot(None)).all()
    agg = aggregate_dsa_summary(dsa_lines)[dsa.id]

    assert agg["wht_amount"] == pytest.approx(agg["total_commission"] * settings.dsa_wht_rate, abs=0.01)
    assert agg["net_salary"] == pytest.approx(agg["total_commission"] - agg["wht_amount"], abs=0.01)


# ---------------------------------------------------------------------
# DSA Summary sheet: headers, SDL/WCF basis (net salary + this run's
# manual adjustments for that DSA), and the informational footnote
# ---------------------------------------------------------------------

def test_dsa_summary_sheet_headers_and_sdl_wcf_with_adjustment(db_session):
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    # A manual clawback applied to this same run, for this DSA.
    db.add(CommissionAdjustment(
        commission_run_id=run.id, payee_type=PayeeType.DSA, dsa_id=dsa.id,
        amount=-1_000.0, reason="test clawback", created_by_user_id=admin.id,
    ))
    db.flush()

    buf = generate_commission_report(db, run)
    wb = openpyxl.load_workbook(buf)
    ws = wb["DSA Summary"]

    header_row = [c.value for c in ws[4]]  # BANNER_ROWS=3 -> header on row 4
    assert header_row == [
        "DSA Code", "DSA Name", "Branch", "DSA Bank Account",
        "New Loans Amount", "Topup Loans Amount",
        "Commission Total (Gross)", "WHT (5%)", "Net Salary",
        "SDL (informational only)", "WCF (informational only)",
    ]

    data_row = [c.value for c in ws[5]]
    gross = round(1_000_000 * settings.dsa_nl_rate, 2)
    wht = round(gross * settings.dsa_wht_rate, 2)
    net_salary = round(gross - wht, 2)
    total_amount = net_salary - 1_000.0  # clawback applied
    expected_sdl = round(total_amount * settings.dsa_sdl_rate, 2)
    expected_wcf = round(total_amount * settings.dsa_wcf_rate, 2)

    assert data_row[6] == gross
    assert data_row[7] == wht
    assert data_row[8] == net_salary
    assert data_row[9] == expected_sdl
    assert data_row[10] == expected_wcf

    # Footnote row (after the data row + totals row) makes clear SDL/WCF
    # are not deducted.
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert SDL_WCF_FOOTNOTE in all_text


def test_dsa_summary_sdl_wcf_without_adjustment_uses_net_salary_only(db_session):
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    buf = generate_commission_report(db, run)
    wb = openpyxl.load_workbook(buf)
    ws = wb["DSA Summary"]
    data_row = [c.value for c in ws[5]]

    gross = round(1_000_000 * settings.dsa_nl_rate, 2)
    wht = round(gross * settings.dsa_wht_rate, 2)
    net_salary = round(gross - wht, 2)
    assert data_row[9] == round(net_salary * settings.dsa_sdl_rate, 2)
    assert data_row[10] == round(net_salary * settings.dsa_wcf_rate, 2)


# ---------------------------------------------------------------------
# Scope guard: DTL Summary report is completely unchanged
# ---------------------------------------------------------------------

def test_dtl_summary_sheet_is_unchanged(db_session):
    db = db_session
    branch, dsa, dtl, admin, branch_upload, biz_upload = _seed_base(db)
    _seed_one_nl_sale(db, branch, branch_upload, biz_upload, disbursement_amt=1_000_000)
    run = _calculate(db, admin)

    buf = generate_commission_report(db, run)
    wb = openpyxl.load_workbook(buf)
    ws = wb["DTL Summary"]
    header_row = [c.value for c in ws[4]]
    assert header_row == ["DTL Name", "Branch", "DSAs Supervised", "DTL Bank Account", "New Loans Amount", "Topup Loans Amount", "Total Commission"]

    data_row = [c.value for c in ws[5]]
    assert data_row[6] == round(1_000_000 * settings.dtl_nl_rate, 2)  # untouched math
