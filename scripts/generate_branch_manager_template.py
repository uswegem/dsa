"""
Regenerates the Branch Manager upload template - the .xlsx handed to a
Branch Manager to fill in and upload (Uploads > Upload Branch Manager
File). Served as a static download from the app itself:
frontend/static/templates/branch_manager_upload_template.xlsx (linked from
the Uploads page - see #bm-template-link in frontend/static/index.html).

Column set/order must match app/services/ingest_branch_manager.py's
EXPECTED_COLUMNS exactly - this script re-parses its own output through the
real ingestion function before saving, so any drift between the two fails
loudly here rather than shipping a template that doesn't actually upload
clean.

Usage (re-run whenever EXPECTED_COLUMNS changes):
    venv/Scripts/python.exe scripts/generate_branch_manager_template.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.ingest_branch_manager import parse_branch_manager_file

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "frontend" / "static" / "templates" / "branch_manager_upload_template.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REQUIRED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

# (header text, note, required?) - header text must resolve to the matching
# field in EXPECTED_COLUMNS (see the parse-back check at the bottom).
COLUMNS = [
    ("DATE", "Date of this loan. Must be a real, parseable date - required on every row.\n\nNote: this is informational only. The commission period is chosen when you upload the file (month/year), not read from this column - a DATE outside that period will be rejected.", True),
    ("CLIENT NAME", "Client's full name.", False),
    ("CLIENT_CHECK_NO", "Client's check/employee number.\nUsed as a secondary match against the Business Manager file - required on every row.", True),
    ("CLIENT_ACCOUNT_NO", "Client's account number.\nThis is the PRIMARY match key against the Business Manager file - required on every row.", True),
    ("APPLICATION NUMBER / ESS", "Application number / ESS reference, if you have one.", False),
    ("REPORTED_AMOUNT", "Reported loan amount (informational - commission is calculated from the Business Manager file's own amounts, not this column).", False),
    ("BRANCH", "Branch name (for reference only). Your branch is applied automatically from your account when you upload - this column is not used to route the data.", False),
    ("DSA_CODE", "The DSA's code - required on every row. Must be unique to that DSA.", True),
    ("DSA_ACCOUNT_NO", "DSA's bank account number, if known.", False),
    ("DSA_NAME", "DSA's full name - required on every row.", True),
    ("DTL_NAME", "Supervising DTL's name, if known.", False),
    ("DTL_CODE", "Supervising DTL's code, if known. Leave blank if you don't have it yet - it can be added later; don't invent one.", False),
]


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    for col, (header, note, required) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = REQUIRED_FILL if required else HEADER_FILL
        cell.font = Font(bold=True) if required else HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.comment = Comment(("REQUIRED. " if required else "Optional. ") + note, "DSA/DTL Commission App")
        ws.column_dimensions[get_column_letter(col)].width = max(len(header) + 4, 16)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return wb


def main():
    wb = build_workbook()

    # Round-trip through the real parser before writing the real output -
    # must resolve every expected column with zero rows/issues on an
    # otherwise-empty template.
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        wb.save(tmp.name)
        result = parse_branch_manager_file(tmp.name, period_month=1, period_year=2000)
    assert result.rows == [], f"template unexpectedly parsed data rows: {result.rows}"
    assert result.issues == [], f"template's headers don't match ingest_branch_manager.EXPECTED_COLUMNS: {result.issues}"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} (verified against parse_branch_manager_file - 0 rows, 0 issues)")


if __name__ == "__main__":
    main()
